"""Build the example statement template: docs/reference/portfolio/statement-template.xlsx.

    python scripts/build_portfolio_statement_template.py [output.xlsx]

The workbook is the template of ``docs/reference/portfolio/statement-template.md``
filled with the demo positions and FICTITIOUS holder data (nobody's real name,
CPF or account). The values are plausible, not real. ``tests/test_portfolio_*``
read this file back through ``src.portfolio.statement``.
"""

from __future__ import annotations

import datetime as dt
import sys
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
DEFAULT_OUT = ROOT / "docs" / "reference" / "portfolio" / "statement-template.xlsx"

POSITION_DATE = dt.date(2026, 9, 30)
TIPOS = ["ação", "fundo", "FII", "ETF", "FIDC", "tesouro", "debênture", "CRI", "CRA", "CDB", "LCI", "LCA", "outro"]
HOLDER = {
    "titular": "MARIA FICTÍCIA DA SILVA",
    "cpf": "123.456.789-09",
    "conta": "12345-6",
    "corretora": "Corretora Exemplo",
}
# (linha_extrato, tipo, codigo, quantidade, preco_unitario[, vencimento, taxa])
DEMO_POSITIONS = [
    ("NTN-B 2035", "tesouro", "NTN-B 2035-05-15", "100", "4123.51"),
    ("PETROBRAS PN", "ação", "PETR4", "20000", "49.40"),
    ("GERAÇÃO L. PAR FIA", "fundo", None, "1500", "176.41"),
    ("XP BANCOS FIC", "fundo", None, "900000", "1.542011"),
    ("XP LIQUIDEZ FIC", "fundo", "51.488.342/0001-33", "700000", "1.289417"),
    ("MN I FIDC", "FIDC", "32113885000121", "750000", "1.00"),
    (
        "BB RENDA FIXA CURTO PRAZO AUTOMÁTICO FUNDO DE INVESTIMENTO EM COTAS DE FUNDOS DE INVESTIMENTO",
        "fundo",
        None,
        "500000",
        "1.541233",
    ),
    ("HGLG11 CSHG LOGISTICA FII", "FII", "HGLG11", "1200", "146.20"),
    # direct credit with the optional columns (engine 1.7): fictitious issuers and registry codes
    ("CDB BANCO EXEMPLO", "CDB", "CDB-26A00001", "1", "180000.00", dt.date(2028, 3, 15), "105,00% do CDI"),
    ("LCA BANCO EXEMPLO", "LCA", "LCA-26A00002", "1", "120000.00", dt.date(2027, 6, 15), "CDI + 0,50%"),
    ("CRA AGRO EXEMPLO", "CRA", "CRA-0260000X", "1", "95000.00", dt.date(2031, 10, 15), "IPCA + 8,74%"),
    ("DEB ENERGIA EXEMPLO", "debênture", "DEB-EXEM12", "1", "60000.00", dt.date(2030, 9, 15), "15,41% a.a."),
    # the same CDB again, as a consolidated statement lists it once per account: merged into one position
    ("CDB BANCO EXEMPLO", "CDB", "CDB-26A00001", "1", "90000.00", dt.date(2028, 3, 15), "105,00% do CDI"),
]


def _valor(qtd: str, preco: str) -> Decimal:
    return (Decimal(qtd) * Decimal(preco)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def build(path: Path = DEFAULT_OUT) -> Decimal:
    wb = Workbook()
    ws = wb.active
    ws.title = "extrato"
    total = sum((_valor(row[3], row[4]) for row in DEMO_POSITIONS), Decimal("0"))
    header_rows = [
        ("titular", HOLDER["titular"]),
        ("cpf", HOLDER["cpf"]),
        ("conta", HOLDER["conta"]),
        ("corretora", HOLDER["corretora"]),
        ("total_extrato", float(total)),
    ]
    for key, value in header_rows:
        ws.append([key, value])
    ws.append([])
    columns = ["linha_extrato", "tipo", "codigo", "quantidade", "preco_unitario", "valor", "data_posicao", "vencimento", "taxa"]
    ws.append(columns)
    header_row = ws.max_row
    for cell in ws[header_row]:
        cell.font = Font(bold=True)
    for name, tipo, codigo, qtd, preco, *opt in DEMO_POSITIONS:
        vencimento, taxa = (opt + [None, None])[:2]
        ws.append([name, tipo, codigo, float(qtd), float(preco), float(_valor(qtd, preco)), POSITION_DATE, vencimento, taxa])
    last = ws.max_row
    for row in ws.iter_rows(min_row=header_row + 1, max_row=last):
        row[6].number_format = "yyyy-mm-dd"
        row[7].number_format = "yyyy-mm-dd"
    dv = DataValidation(type="list", formula1='"' + ",".join(TIPOS) + '"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(f"B{header_row + 1}:B{max(last, header_row + 200)}")
    widths = {"A": 46, "B": 14, "C": 22, "D": 14, "E": 16, "F": 16, "G": 14, "H": 14, "I": 18}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return total


if __name__ == "__main__":
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else DEFAULT_OUT
    print(f"wrote {out} (total_extrato R$ {build(out)})")
