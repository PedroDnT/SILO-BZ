"""Offline tests of the engine's HTTP server (src/portfolio/server.py).

SILO is read through the FakeClient over the canned rows, as in
test_portfolio_engine.py, and the report runs with SILO_LLM_PROVIDER=fake: no
network, no key. The PDF renderer is a stub passed to ``create_app``; the real
one runs only where WeasyPrint is installed (the engine image's CI job runs this
file inside the image). These tests cover the HTTP mapping; the diagnosis itself
(engine, report, trace record) is tested through ``diagnosis.diagnose`` in
test_portfolio_diagnosis.py.
"""

from __future__ import annotations

import io
import logging
from pathlib import Path

import pytest

pytest.importorskip("flask")
pytest.importorskip("pydantic")

from src.portfolio import server  # noqa: E402
from src.portfolio.client import FakeClient, load_fake_rows  # noqa: E402

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / "docs" / "reference" / "portfolio" / "statement-template.xlsx"
FAKE_ROWS = ROOT / "tests" / "fixtures" / "portfolio" / "fake_silo_rows.json"
TOKEN = "test-token-not-a-secret"
UPLOAD_NAME = "MARIA_FICTICIA_12345-6.xlsx"

# What the template carries about the holder; none of it may reach a log line.
PRIVATE = ["MARIA FICTÍCIA", "MARIA", "FICTICIA", "123.456.789-09", "12345678909", "12345-6", UPLOAD_NAME]


def _client():
    return FakeClient(load_fake_rows(FAKE_ROWS))


@pytest.fixture()
def app(monkeypatch):
    monkeypatch.setenv(server.TOKEN_ENV, TOKEN)
    monkeypatch.setenv("SILO_LLM_PROVIDER", "fake")
    return server.create_app(client_factory=_client, pdf_renderer=stub_pdf).test_client()


def _auth(token: str = TOKEN) -> dict:
    return {"Authorization": f"Bearer {token}"}


def stub_pdf(html_text, out_path):
    assert "<html" in html_text.lower()
    Path(out_path).write_bytes(b"%PDF-1.7\n% stub\n")
    return Path(out_path)


def _assert_no_private(caplog, capfd):
    out, err = capfd.readouterr()
    text = caplog.text + out + err
    for s in PRIVATE:
        assert s not in text, f"{s!r} reached the logs"


def test_health(app):
    r = app.get("/health")
    assert r.status_code == 200 and r.data == b"ok\n"


def test_every_answer_carries_the_engine_revision(app):
    rev = server.engine_rev()
    assert len(rev) == 12 and int(rev, 16) >= 0
    assert app.get("/health").headers["X-Silo-Engine-Rev"] == rev
    assert app.post("/diagnose", data=b"x").headers["X-Silo-Engine-Rev"] == rev


def test_engine_rev_ignores_bytecode_and_follows_content(tmp_path):
    (tmp_path / "a.py").write_text("x = 1\n")
    (tmp_path / "__pycache__").mkdir()
    (tmp_path / "__pycache__" / "a.cpython-312.pyc").write_bytes(b"junk")
    first = server.engine_rev(tmp_path)
    (tmp_path / "__pycache__" / "a.cpython-312.pyc").write_bytes(b"other")
    assert server.engine_rev(tmp_path) == first
    (tmp_path / "a.py").write_text("x = 2\n")
    assert server.engine_rev(tmp_path) != first


def test_health_answers_without_the_token_configured(monkeypatch):
    monkeypatch.delenv(server.TOKEN_ENV, raising=False)
    r = server.create_app(client_factory=_client).test_client().get("/health")
    assert r.status_code == 200


@pytest.mark.parametrize("headers", [{}, _auth("wrong"), {"Authorization": TOKEN}, {"Authorization": f"Basic {TOKEN}"}])
def test_diagnose_refuses_without_the_right_bearer_token(app, headers):
    r = app.post("/diagnose", data=TEMPLATE.read_bytes(), headers=headers)
    assert r.status_code == 401
    assert r.json == {"erro": server.MSG[401]}


@pytest.mark.parametrize("value", [None, ""])
def test_diagnose_refuses_everything_when_the_token_is_not_configured(monkeypatch, value):
    if value is None:
        monkeypatch.delenv(server.TOKEN_ENV, raising=False)
    else:
        monkeypatch.setenv(server.TOKEN_ENV, value)
    c = server.create_app(client_factory=_client).test_client()
    for headers in ({}, _auth(""), _auth(TOKEN)):
        r = c.post("/diagnose", data=b"x", headers=headers)
        assert r.status_code == 503 and r.json == {"erro": server.MSG[503]}


def test_upload_over_the_limit_is_refused(app):
    big = b"PK\x03\x04" + b"0" * server.MAX_UPLOAD_BYTES
    r = app.post("/diagnose", data=big, headers=_auth())
    assert r.status_code == 413 and r.json == {"erro": server.MSG[413]}
    r = app.post(
        "/diagnose",
        data={"file": (io.BytesIO(big), UPLOAD_NAME)},
        content_type="multipart/form-data",
        headers=_auth(),
    )
    assert r.status_code == 413


def test_empty_and_unknown_uploads(app):
    assert app.post("/diagnose", data=b"", headers=_auth()).status_code == 400
    r = app.post("/diagnose", data=b"titular;cpf\n", headers=_auth())
    assert r.status_code == 415 and r.json == {"erro": server.MSG[415]}


def test_unreadable_statement_is_a_fixed_422_with_nothing_logged(app, caplog, capfd):
    caplog.set_level(logging.DEBUG)
    r = app.post("/diagnose", data=b"PK\x03\x04 " + "MARIA FICTÍCIA 123.456.789-09 not a zip".encode(), headers=_auth())
    assert r.status_code == 422 and r.json == {"erro": server.MSG[422]}
    _assert_no_private(caplog, capfd)


def test_statement_that_does_not_reconcile_is_a_fixed_422(app, caplog, capfd, tmp_path):
    from openpyxl import load_workbook

    wb = load_workbook(TEMPLATE)
    ws = wb.worksheets[0]
    for row in ws.iter_rows():
        if row[0].value == "total_extrato":
            row[1].value = 1.0
    bad = tmp_path / "bad.xlsx"
    wb.save(bad)
    caplog.set_level(logging.DEBUG)
    r = app.post(
        "/diagnose",
        data={"file": (io.BytesIO(bad.read_bytes()), UPLOAD_NAME)},
        content_type="multipart/form-data",
        headers=_auth(),
    )
    assert r.status_code == 422 and r.json == {"erro": server.MSG[422]}
    assert b"5651424" not in r.data and b"MARIA" not in r.data
    _assert_no_private(caplog, capfd)


def test_engine_failure_is_a_fixed_500_with_no_traceback(monkeypatch, caplog, capfd):
    monkeypatch.setenv(server.TOKEN_ENV, TOKEN)

    def boom():
        raise RuntimeError("MARIA FICTÍCIA 12345-6")

    c = server.create_app(client_factory=boom).test_client()
    caplog.set_level(logging.DEBUG)
    r = c.post("/diagnose", data=TEMPLATE.read_bytes(), headers=_auth())
    assert r.status_code == 500 and r.json == {"erro": server.MSG[500]}
    assert r.headers["X-Silo-Stage"] == "engine"
    assert r.headers["X-Silo-Error"] == "RuntimeError" and "X-Silo-Error-Status" not in r.headers
    assert "MARIA" not in str(r.headers)
    assert "Traceback" not in caplog.text
    _assert_no_private(caplog, capfd)


def test_provider_error_code_is_named_but_never_free_text(monkeypatch):
    monkeypatch.setenv(server.TOKEN_ENV, TOKEN)

    class FakeApiError(Exception):
        status_code = 403
        code = "unsupported_country_region_territory"

    class FreeText(Exception):
        status_code = 400
        code = "MARIA FICTÍCIA tem 12345-6"

    for exc, want in ((FakeApiError("x"), "unsupported_country_region_territory"), (FreeText("x"), None)):
        def boom(exc=exc):
            raise exc

        r = server.create_app(client_factory=boom).test_client().post("/diagnose", data=TEMPLATE.read_bytes(), headers=_auth())
        assert r.status_code == 500
        assert r.headers.get("X-Silo-Error-Code") == want
        assert "MARIA" not in str(r.headers)


def test_missing_llm_key_is_a_fixed_502(monkeypatch, app):
    monkeypatch.setenv("SILO_LLM_PROVIDER", "anthropic")
    monkeypatch.delenv("ANTHROPIC_API_KEY", raising=False)
    r = app.post("/diagnose", data=TEMPLATE.read_bytes(), headers=_auth())
    assert r.status_code == 502 and r.json == {"erro": server.MSG[502]}
    assert r.headers["X-Silo-Stage"] == "report"


@pytest.mark.parametrize("multipart", [False, True])
def test_diagnose_returns_a_pdf_with_the_renderer_stubbed(app, caplog, capfd, multipart):
    caplog.set_level(logging.DEBUG)
    if multipart:
        r = app.post(
            "/diagnose",
            data={"file": (io.BytesIO(TEMPLATE.read_bytes()), UPLOAD_NAME)},
            content_type="multipart/form-data",
            headers=_auth(),
        )
    else:
        r = app.post("/diagnose", data=TEMPLATE.read_bytes(), headers=_auth())
    assert r.status_code == 200, r.data[:200]
    assert r.mimetype == "application/pdf" and r.data.startswith(b"%PDF")
    assert r.headers["Cache-Control"] == "no-store"
    assert r.headers["X-Silo-Provider"] == "fake"
    assert r.headers["X-Silo-Narrative"] and float(r.headers["X-Silo-Cost-Usd"]) >= 0
    assert float(r.headers["X-Silo-Seconds"]) >= 0
    assert "diagnose 200 format=xlsx" in caplog.text
    _assert_no_private(caplog, capfd)


def test_diagnose_renders_a_real_pdf(monkeypatch, caplog, capfd):
    pytest.importorskip("weasyprint")
    monkeypatch.setenv(server.TOKEN_ENV, TOKEN)
    monkeypatch.setenv("SILO_LLM_PROVIDER", "fake")
    app = server.create_app(client_factory=_client).test_client()  # the default renderer: html_to_pdf
    caplog.set_level(logging.DEBUG)
    r = app.post(
        "/diagnose",
        data={"file": (io.BytesIO(TEMPLATE.read_bytes()), UPLOAD_NAME)},
        content_type="multipart/form-data",
        headers=_auth(),
    )
    assert r.status_code == 200, r.data[:200]
    assert r.mimetype == "application/pdf" and r.data.startswith(b"%PDF")
    assert len(r.data) > 10_000
    _assert_no_private(caplog, capfd)


def test_temporary_files_are_removed(app, monkeypatch, tmp_path):
    monkeypatch.setattr(server.tempfile, "tempdir", str(tmp_path))
    assert app.post("/diagnose", data=TEMPLATE.read_bytes(), headers=_auth()).status_code == 200
    assert app.post("/diagnose", data=b"PK\x03\x04 broken", headers=_auth()).status_code == 422
    assert list(tmp_path.iterdir()) == []


def test_unknown_path_and_method_answer_json(app):
    assert app.get("/nope").status_code == 404
    r = app.get("/diagnose")
    assert r.status_code == 405 and r.json == {"erro": server.MSG[405]}


def test_narrative_headers_name_the_reason_and_the_calls_never_free_text():
    from src.portfolio.report.render import Narrative

    n = Narrative(status="unknown", reason="LLMOutputError: o modelo escreveu MARIA", reason_code="LLMOutputError")
    n.calls = [{"role": "redator", "model": "gpt-5.1-2025-11-13", "cost_usd": 0.17, "output_tokens": 16000,
                "reasoning_tokens": 15800, "input_tokens": 9000}]
    h = server.narrative_headers(n)
    assert h == {"X-Silo-Narrative-Reason": "LLMOutputError",
                 "X-Silo-Llm-Calls": "redator:out=16000:reasoning=15800"}
    assert "MARIA" not in str(h)
    # A code that is not a bare identifier is dropped, never echoed.
    assert server.narrative_headers(Narrative(status="unknown", reason_code="x: MARIA")) == {}
    assert server.narrative_headers(Narrative(status="complete")) == {}


def test_html_delivery_skips_pdf_and_has_client_constraints(monkeypatch):
    def forbidden_pdf(*args):
        raise AssertionError('HTML must not render a PDF')
    monkeypatch.setenv(server.TOKEN_ENV, TOKEN)
    monkeypatch.setenv("SILO_LLM_PROVIDER", "fake")
    app = server.create_app(client_factory=_client, pdf_renderer=forbidden_pdf).test_client()
    r = app.post('/diagnose', headers=_auth(), data={
        'file': (io.BytesIO(TEMPLATE.read_bytes()), 'statement.xlsx'),
        'output_format': 'html',
        'client_constraints': '{"profile":"conservador","horizon_date":"2027-01-01","liquidity_brl":"1000","liquidity_date":"2027-01-01"}',
    })
    assert r.status_code == 200 and r.mimetype == 'text/html'
    assert b'Resumo para a reuni' in r.data
    assert b'<details id="apendice">' in r.data
    assert b'conservador' in r.data and b'suitability' in r.data
    assert r.headers['Cache-Control'] == 'no-store'
    assert "default-src 'none'" in r.headers['Content-Security-Policy']


def test_client_constraints_refused_before_engine(monkeypatch):
    def forbidden_client():
        raise AssertionError('invalid input must not reach SILO')
    monkeypatch.setenv(server.TOKEN_ENV, TOKEN)
    app = server.create_app(client_factory=forbidden_client).test_client()
    r = app.post('/diagnose', headers=_auth(), data={
        'file': (io.BytesIO(TEMPLATE.read_bytes()), 'statement.xlsx'),
        'client_constraints': '{"horizon_date":"2000-01-01"}',
    })
    assert r.status_code == 400
    assert r.headers['X-Silo-Stage'] == 'engine'


@pytest.mark.parametrize('bad', ['{"horizon_date":"2000-01-01"}', '{"profile":"temerario"}', '{"other":1}',
                                 '["not","a","dict"]', '{"liquidity_brl":"100"}', 'not json'])
def test_every_bad_constraint_is_the_same_400(monkeypatch, bad):
    def forbidden_client():
        raise AssertionError('invalid input must not reach SILO')
    monkeypatch.setenv(server.TOKEN_ENV, TOKEN)
    app = server.create_app(client_factory=forbidden_client).test_client()
    r = app.post('/diagnose', headers=_auth(), data={
        'file': (io.BytesIO(TEMPLATE.read_bytes()), 'statement.xlsx'), 'client_constraints': bad})
    assert r.status_code == 400


def test_curl_default_form_content_type_does_not_consume_raw_upload(app):
    # Exact smoke wire format: curl --data-binary sends application/x-www-form-urlencoded.
    headers = {**_auth(), 'Content-Type': 'application/x-www-form-urlencoded'}
    r = app.post('/diagnose', data=b'not a statement', headers=headers)
    assert r.status_code == 415
    assert app.post('/diagnose', data=b'', headers=headers).status_code == 400
    r = app.post('/diagnose', data=TEMPLATE.read_bytes(), headers=headers)
    assert r.status_code == 200 and r.mimetype == 'application/pdf'
