"""Reader for the broker-statement spreadsheet template (xlsx or csv).

Spec: docs/reference/portfolio/statement-template.md. One sheet (the first);
a block of ``key | value`` header rows (``titular``, ``cpf``, ``conta``,
``corretora``, ``total_extrato``), then a header row naming the position
columns, then one row per position.

Two hard rules:

* Masking happens here, before anything else sees the data: the holder fields
  build a ``Masker`` that scrubs every string cell, and the ``Statement``
  returned carries only tokens (``src/portfolio/mask.py``).
* The sum of ``valor`` must equal ``total_extrato`` within R$0.01 times the
  number of position rows, and every position row must be readable. Otherwise
  ``StatementTotalMismatch`` names the difference and every row it could not
  read. No row is ever dropped silently.
"""

from __future__ import annotations

import csv
import datetime as dt
import io
import re
import unicodedata
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path

from src.portfolio.mask import MaskedHolder, Masker

POSITION_COLUMNS = (
    "linha_extrato",
    "tipo",
    "codigo",
    "quantidade",
    "preco_unitario",
    "valor",
    "data_posicao",
    # optional (engine 1.7): the maturity and the rate exactly as the statement prints them
    "vencimento",
    "taxa",
)
REQUIRED_COLUMNS = ("linha_extrato", "tipo", "valor", "data_posicao")
HOLDER_KEYS = ("titular", "cpf", "conta")
HEADER_KEYS = HOLDER_KEYS + ("corretora", "total_extrato")

# Canonical spelling of each statement type, keyed by its accent-free lower form.
TIPOS = {
    "acao": "ação",
    "fundo": "fundo",
    "fii": "FII",
    "etf": "ETF",
    "fidc": "FIDC",
    "tesouro": "tesouro",
    "debenture": "debênture",
    "cri": "CRI",
    "cra": "CRA",
    "cdb": "CDB",
    "lci": "LCI",
    "lca": "LCA",
    "outro": "outro",
}

CENT = Decimal("0.01")


class StatementError(ValueError):
    """Base class: the statement cannot be read as a whole."""


class StatementFormatError(StatementError):
    """The file is not the template (no header row, no total, unknown layout)."""


@dataclass(frozen=True)
class UnreadableRow:
    source_row: int
    reason: str

    def as_dict(self) -> dict[str, object]:
        return {"source_row": self.source_row, "reason": self.reason}


class StatementTotalMismatch(StatementError):
    """The positions do not add up to the statement's own total, or a row could not be read."""

    def __init__(
        self,
        *,
        stated_total: Decimal,
        sum_of_lines: Decimal,
        tolerance: Decimal,
        n_rows: int,
        unreadable_rows: list[UnreadableRow],
        check: str | None = None,
        failures: list[str] | None = None,
    ):
        self.check = check
        self.failures = list(failures or [])
        self.stated_total = stated_total
        self.sum_of_lines = sum_of_lines
        self.difference = sum_of_lines - stated_total
        self.tolerance = tolerance
        self.n_rows = n_rows
        self.unreadable_rows = list(unreadable_rows)
        rows = "; ".join(f"row {u.source_row}: {u.reason}" for u in self.unreadable_rows) or "none"
        head = f"[{check}] " if check else ""
        extra = (" Checks failed: " + " | ".join(self.failures) + ".") if self.failures else ""
        super().__init__(
            f"{head}statement does not reconcile: sum of readable lines R$ {sum_of_lines} vs stated total "
            f"R$ {stated_total}, difference R$ {self.difference} (tolerance R$ {tolerance} = "
            f"R$ 0.01 x {n_rows} rows). Rows not read: {rows}.{extra}"
        )


@dataclass(frozen=True)
class Position:
    line_no: int  # 1-based order among position rows
    source_row: int  # 1-based row number in the sheet / csv
    linha_extrato: str
    tipo: str
    codigo: str | None
    quantidade: Decimal | None
    preco_unitario: Decimal | None
    valor: Decimal
    data_posicao: dt.date
    # Optional: the PDF reader fills them, and the spreadsheet reader from its optional columns 'vencimento' and 'taxa'.
    vencimento: dt.date | None = None
    taxa_texto: str | None = None  # the rate exactly as the statement prints it
    estrategia_corretora: str | None = None  # the broker's own labels, never ours
    classe_corretora: str | None = None
    conta_ref: str | None = None  # an ordinal token (C1, C2...), never the real account
    preco_implicito: bool = False  # preco_unitario was derived as valor / quantidade
    contas: tuple["ContaLine", ...] = ()  # per-account lines of a consolidated position
    emissor: str | None = None  # the issuer as the statement prints it (BTG extrato), for issuer concentration


@dataclass(frozen=True)
class ContaLine:
    """One account's share of a consolidated position."""

    conta_ref: str | None  # None when the lines merged came from one statement (no account is invented)
    titular_ref: str | None
    source_row: int
    linha_extrato: str
    quantidade: Decimal | None
    preco_unitario: Decimal | None
    valor: Decimal
    data_posicao: dt.date


@dataclass(frozen=True)
class AccountSummary:
    conta_ref: str
    titular_ref: str | None
    n_lines: int
    stated_total: Decimal
    sum_of_lines: Decimal
    position_date: dt.date
    source_format: str
    positions: tuple["Position", ...] = ()  # this account's own positions (conta_ref set), the per-account view


@dataclass(frozen=True)
class Statement:
    holder: MaskedHolder
    corretora: str | None
    stated_total: Decimal
    sum_of_lines: Decimal
    tolerance: Decimal
    positions: tuple[Position, ...]
    position_date: dt.date
    position_dates: tuple[dt.date, ...]
    source_format: str  # "xlsx", "csv" or "pdf"; the file name is not kept (it may name the client)
    notes: tuple[str, ...] = field(default_factory=tuple)
    accounts: tuple[AccountSummary, ...] = ()  # one per consolidated statement; empty for a single file


# ---------------------------------------------------------------------------
# Cell parsing.
# ---------------------------------------------------------------------------


def _norm_key(value: object) -> str:
    s = "" if value is None else str(value)
    s = "".join(c for c in unicodedata.normalize("NFKD", s) if not unicodedata.combining(c))
    return s.strip().lower()


def parse_decimal(value: object) -> Decimal | None:
    """A numeric cell, or text in pt-BR (``1.234,56``) or plain (``1234.56``) form."""
    if value is None:
        return None
    if isinstance(value, bool):
        raise ValueError("boolean is not a number")
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value))
    s = str(value).strip().replace("R$", "").replace(" ", "").replace(" ", "")
    if s == "":
        return None
    if "," in s:
        s = s.replace(".", "").replace(",", ".")
    try:
        return Decimal(s)
    except InvalidOperation as exc:
        raise ValueError("not a number") from exc


def parse_date(value: object) -> dt.date | None:
    if value is None:
        return None
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    s = str(value).strip()
    if s == "":
        return None
    for fmt in ("%Y-%m-%d", "%d/%m/%Y"):
        try:
            return dt.datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    raise ValueError("not a date (use YYYY-MM-DD or DD/MM/YYYY)")


def _text(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    s = str(value).strip()
    return s or None


# ---------------------------------------------------------------------------
# File reading: rows of raw cells. Nothing here is logged.
# ---------------------------------------------------------------------------


def _rows_from_xlsx(path: Path) -> list[list[object]]:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        return [list(r) for r in ws.iter_rows(values_only=True)]
    finally:
        wb.close()


def _rows_from_csv(path: Path) -> list[list[object]]:
    raw = path.read_bytes()
    text = None
    for enc in ("utf-8-sig", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    assert text is not None  # latin-1 decodes any byte string
    first = text.splitlines()[0] if text else ""
    delim = ";" if first.count(";") > first.count(",") else ","
    return [list(r) for r in csv.reader(io.StringIO(text), delimiter=delim)]


def read_statement(path: str | Path) -> Statement:
    """Read, mask and reconcile one statement file."""
    p = Path(path)
    suffix = p.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        rows = _rows_from_xlsx(p)
        fmt = "xlsx"
    elif suffix in (".csv", ".txt"):
        rows = _rows_from_csv(p)
        fmt = "csv"
    else:
        raise StatementFormatError(f"unsupported statement format {suffix!r}; use .xlsx or .csv")
    return parse_rows(rows, source_format=fmt)


def _is_blank(row: list[object]) -> bool:
    return all(_text(c) is None for c in row)


def parse_rows(rows: list[list[object]], *, source_format: str = "rows") -> Statement:
    """Parse the template from raw rows (the first sheet, or the csv lines)."""
    # 1. Locate the position header row; everything above is the header block.
    header_idx = None
    for i, row in enumerate(rows):
        if any(_norm_key(c) == "linha_extrato" for c in row):
            header_idx = i
            break
    if header_idx is None:
        raise StatementFormatError("no position header row: a row naming the column 'linha_extrato' is required")

    # 2. Header block: key in the first non-empty cell, value in the next one.
    header: dict[str, object] = {}
    for row in rows[:header_idx]:
        cells = [c for c in row if _text(c) is not None]
        if not cells:
            continue
        key = _norm_key(cells[0])
        if key in HEADER_KEYS:
            # Never echo a header value in an error: it may be personal data.
            if key in header:
                raise StatementFormatError(f"header key '{key}' appears twice")
            header[key] = cells[1] if len(cells) > 1 else None

    masker = Masker(header.get("titular"), header.get("cpf"), header.get("conta"))
    try:
        return _parse_body(rows, header_idx, header, masker, source_format)
    finally:
        del masker


def _parse_body(
    rows: list[list[object]],
    header_idx: int,
    header: dict[str, object],
    masker: Masker,
    source_format: str,
) -> Statement:
    if "total_extrato" not in header or _text(header["total_extrato"]) is None:
        raise StatementFormatError("header cell 'total_extrato' (the statement's own total) is required")
    try:
        stated_total = parse_decimal(header["total_extrato"])
    except ValueError:
        raise StatementFormatError("header cell 'total_extrato' is not a number") from None
    assert stated_total is not None

    columns: dict[str, int] = {}
    for j, cell in enumerate(rows[header_idx]):
        key = _norm_key(cell)
        if key in POSITION_COLUMNS:
            columns[key] = j
    missing = [c for c in REQUIRED_COLUMNS if c not in columns]
    if missing:
        raise StatementFormatError(f"position header is missing column(s): {', '.join(missing)}")

    def cell(row: list[object], name: str) -> object:
        j = columns.get(name)
        if j is None or j >= len(row):
            return None
        return masker.scrub(row[j])

    positions: list[Position] = []
    unreadable: list[UnreadableRow] = []
    n_rows = 0
    for i in range(header_idx + 1, len(rows)):
        row = rows[i]
        if _is_blank(row):
            continue
        n_rows += 1
        source_row = i + 1
        problems: list[str] = []
        linha = _text(cell(row, "linha_extrato"))
        if linha is None:
            problems.append("linha_extrato empty")
        tipo_raw = _text(cell(row, "tipo"))
        tipo = TIPOS.get(_norm_key(tipo_raw)) if tipo_raw else None
        if tipo is None:
            problems.append("tipo empty or not one of the template's types")
        try:
            valor = parse_decimal(cell(row, "valor"))
            if valor is None:
                problems.append("valor empty")
        except ValueError:
            valor = None
            problems.append("valor is not a number")
        try:
            qtd = parse_decimal(cell(row, "quantidade"))
        except ValueError:
            qtd = None
            problems.append("quantidade is not a number")
        try:
            preco = parse_decimal(cell(row, "preco_unitario"))
        except ValueError:
            preco = None
            problems.append("preco_unitario is not a number")
        try:
            data = parse_date(cell(row, "data_posicao"))
            if data is None:
                problems.append("data_posicao empty")
        except ValueError:
            data = None
            problems.append("data_posicao is not a date")
        try:
            vencimento = parse_date(cell(row, "vencimento"))
        except ValueError:
            vencimento = None
            problems.append("vencimento is not a date")
        if problems:
            unreadable.append(UnreadableRow(source_row, "; ".join(problems)))
            continue
        assert linha is not None and tipo is not None and valor is not None and data is not None
        positions.append(
            Position(
                line_no=len(positions) + 1,
                source_row=source_row,
                linha_extrato=linha,
                tipo=tipo,
                codigo=_text(cell(row, "codigo")),
                quantidade=qtd,
                preco_unitario=preco,
                valor=valor,
                data_posicao=data,
                vencimento=vencimento,
                taxa_texto=_text(cell(row, "taxa")),
            )
        )

    if n_rows == 0:
        raise StatementFormatError("the statement has no position rows")
    sum_of_lines = sum((p.valor for p in positions), Decimal("0"))
    tolerance = CENT * n_rows
    if unreadable or abs(sum_of_lines - stated_total) > tolerance:
        raise StatementTotalMismatch(
            stated_total=stated_total,
            sum_of_lines=sum_of_lines,
            tolerance=tolerance,
            n_rows=n_rows,
            unreadable_rows=unreadable,
        )

    dates = tuple(sorted({p.data_posicao for p in positions}))
    notes: list[str] = []
    if len(dates) > 1:
        notes.append(
            "As linhas têm datas de posição diferentes; a data de referência do relatório é a mais recente "
            f"({dates[-1].isoformat()})."
        )
    return Statement(
        holder=masker.holder,
        corretora=_text(masker.scrub(_text(header.get("corretora")))),
        stated_total=stated_total,
        sum_of_lines=sum_of_lines,
        tolerance=tolerance,
        positions=tuple(positions),
        position_date=dates[-1],
        position_dates=dates,
        source_format=source_format,
        notes=tuple(notes),
    )


_CNPJ_RE = re.compile(r"^\d{14}$")


def codigo_cnpj(codigo: str | None) -> str | None:
    """The 14-digit CNPJ in a ``codigo`` cell, when the cell is a CNPJ (punctuation allowed)."""
    if not codigo:
        return None
    s = re.sub(r"[.\-/\s]", "", codigo)
    return s if _CNPJ_RE.match(s) else None
